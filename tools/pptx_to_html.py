"""Write an accessible HTML transcript for each public slide deck.

Reads the public copies in slides/<stem>.pptx (made by make_public_decks.ps1;
never the originals, so speaker notes and hidden slides cannot leak) and writes
slides/<stem>.html next to each one. Each page is self-contained: inline CSS, no
JavaScript, no external fonts, scripts, or images.

Page structure:

* a skip link to the main content;
* a header with the deck title, the workshop and event line, the presenters and
  date, and links to the same deck as PowerPoint and PDF and to the repository;
* a navigation list with one link per slide;
* one section per slide with a "Slide N: title" heading, followed by the slide's
  content in reading order: title first, then the remaining shapes top to bottom
  and left to right, recursing into groups. Bulleted and numbered paragraphs
  become nested lists, tables become tables with a header row, hyperlinks are
  kept, and the alt text of images, charts, and diagrams is written out as text.

Slide-number, footer, and date placeholders are skipped. Pictures marked
decorative or without alt text are skipped. Alt text that is only a filename is
treated as missing and reported, as is any chart or diagram without alt text and
any slide without a title.

A picture whose alt text ends with a colon (for example "Instagram:") is a
label for the text beside it, not an image in its own right. It is paired with
the nearest text shape to its right in the same vertical band, and the two are
written as one paragraph, "Instagram: @UM_NaNDA", with any hyperlink on the
text kept and no "Image:" prefix. A label with no text beside it is written as
an image description.

Usage:
    python tools/pptx_to_html.py                 # every deck in DECKS
    python tools/pptx_to_html.py 04_Choosing_Your_First_NaNDA_Dataset

Requires python-pptx (pip install python-pptx).
"""

from __future__ import annotations

import html
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

REPO_ROOT = Path(__file__).resolve().parents[1]
SLIDES_DIR = REPO_ROOT / "slides"
REPO_URL = (
    "https://github.com/The-National-Neighborhood-Data-Archive/"
    "GSA_Neighborhood-Data-in-Aging-Research_2026"
)

WORKSHOP_TITLE = (
    "Using Neighborhood Data in Aging Research: "
    "The National Neighborhood Data Archive (NaNDA)"
)
EVENT_LINE = "GSA Annual Scientific Meeting Online Workshop Series 2026"
DATE_LINE = "Wednesday, October 7, 2026"

# One entry per deck: public filename stem, session name, presenters, and the
# page heading (None means the heading is the session name). The session name
# is the <title> of the page and matches the Title property of the original
# deck, so the transcript, the PDF, and the README all call a session by the
# same name. The talk titles for sessions 2 to 4 are the presenters' own.
DECKS = [
    ("01_Welcome_and_Overview",
     "Welcome and overview: what NaNDA is and a tour of the repository",
     "Lindsay Gypin and Philippa Clarke",
     None),
    ("02_Neighborhood_Environment_and_Aging_with_Disability",
     "Maintaining Health while Aging with Disability: "
     "The Role of the Neighborhood Environment",
     "Philippa Clarke",
     None),
    ("03_NaNDA_and_Immune_Aging",
     "Where We Live Gets Under the Skin: NaNDA and Immune Aging",
     "Grace Noppert",
     None),
    ("04_Choosing_Your_First_NaNDA_Dataset",
     "Choosing Your First NaNDA Dataset",
     "Lindsay Gypin",
     "Breakout 1 (beginner): Choosing Your First NaNDA Dataset"),
    ("05_Reconvene_and_Wrap-Up",
     "Reconvene and wrap-up: report back, resources, and next steps",
     "Grace Noppert and Lindsay Gypin",
     None),
]

SKIP_PLACEHOLDERS = {
    PP_PLACEHOLDER.SLIDE_NUMBER,
    PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.DATE,
}
TITLE_PLACEHOLDERS = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.VERTICAL_TITLE,
}
BODY_LIKE_PLACEHOLDERS = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.SUBTITLE,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
}

# Alt text that is really a filename counts as missing.
FILENAME_ALT = re.compile(
    r"^[\w\-. ()]+\.(png|jpe?g|gif|bmp|tiff?|svg|emf|wmf|webp)$", re.IGNORECASE
)

# Shapes whose top edges are within this distance (EMU, 0.25 inch) are read as
# one row, left to right. This keeps a step number and the text beside it
# together even when the number sits a little lower than the text box.
ROW_TOLERANCE = 228600

NV_PR_TAGS = (
    qn("p:nvSpPr"),
    qn("p:nvPicPr"),
    qn("p:nvGraphicFramePr"),
    qn("p:nvGrpSpPr"),
    qn("p:nvCxnSpPr"),
)

CSS = """
:root { color-scheme: light; }
body {
  margin: 0;
  font-family: system-ui, -apple-system, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
  font-size: 1.05rem;
  line-height: 1.5;
  color: #1a1a1a;
  background: #ffffff;
}
.page { max-width: 72ch; margin: 0 auto; padding: 1.5rem 1rem 4rem; }
a { color: #0b5394; text-decoration: underline; }
a:hover { text-decoration-thickness: 2px; }
a:focus { outline: 3px solid #0b5394; outline-offset: 3px; }
.skip-link {
  position: absolute;
  left: -10000px;
  top: 0;
  padding: 0.5rem 1rem;
  background: #ffffff;
  color: #0b5394;
  border: 3px solid #0b5394;
  font-weight: 600;
}
.skip-link:focus { left: 1rem; top: 1rem; z-index: 1; }
h1 { font-size: 2rem; line-height: 1.2; margin: 0 0 0.75rem; }
h2 { font-size: 1.5rem; line-height: 1.25; margin: 2.5rem 0 0.75rem; }
main h2 { padding-top: 1rem; border-top: 1px solid #cccccc; }
h3 { font-size: 1.2rem; line-height: 1.3; margin: 1.5rem 0 0.5rem; }
p, ul, ol { margin: 0 0 0.75rem; }
ul, ol { padding-left: 1.5rem; }
li { margin-bottom: 0.25rem; }
li > ul, li > ol { margin-top: 0.25rem; }
.about { color: #333333; }
.image { font-style: italic; }
nav ol { columns: 1; }
table { border-collapse: collapse; width: 100%; margin: 0 0 1rem; }
th, td { border: 1px solid #767676; padding: 0.4rem 0.6rem; text-align: left; vertical-align: top; }
th { background: #e8eef5; color: #1a1a1a; }
@media (min-width: 40em) { nav ol { columns: 2; column-gap: 2rem; } }
@media (prefers-reduced-motion: reduce) {
  *, *::before, *::after { animation: none !important; transition: none !important; scroll-behavior: auto !important; }
}
""".strip()


@dataclass
class Report:
    stem: str
    slide_count: int = 0
    missing_alt: list[tuple[int, str, str]] = field(default_factory=list)
    missing_title: list[int] = field(default_factory=list)
    links: set[str] = field(default_factory=set)


# --------------------------------------------------------------------------- #
# Small XML helpers
# --------------------------------------------------------------------------- #

def esc(text: str) -> str:
    return html.escape(text, quote=False)


def attr(text: str) -> str:
    return html.escape(text, quote=True)


def cnvpr(element):
    """Return the p:cNvPr element of a shape element, or None."""
    for child in element:
        if child.tag in NV_PR_TAGS:
            return child.find(qn("p:cNvPr"))
    return None


def alt_text(shape) -> str:
    c = cnvpr(shape._element)
    if c is None:
        return ""
    return (c.get("descr") or "").strip()


def is_decorative(shape) -> bool:
    c = cnvpr(shape._element)
    if c is None:
        return False
    for el in c.iter():
        if el.tag.endswith("}decorative") and el.get("val") in ("1", "true"):
            return True
    return False


def resolve_link(part, hlink) -> str | None:
    """Return the external address of an a:hlinkClick element, or None."""
    if hlink is None:
        return None
    rid = hlink.get(qn("r:id"))
    if not rid:
        return None
    try:
        rel = part.rels[rid]
    except KeyError:
        return None
    if not rel.is_external:
        return None
    return rel.target_ref


def shape_link(shape) -> str | None:
    c = cnvpr(shape._element)
    if c is None:
        return None
    return resolve_link(shape.part, c.find(qn("a:hlinkClick")))


# --------------------------------------------------------------------------- #
# Bullets: explicit paragraph properties first, then the inheritance chain
# (shape list style, layout placeholder, master placeholder, master text styles)
# --------------------------------------------------------------------------- #

def bullet_from_ppr(ppr):
    """Return 'ul', 'ol', None (explicit no bullet), or 'inherit'."""
    if ppr is None:
        return "inherit"
    if ppr.find(qn("a:buNone")) is not None:
        return None
    if ppr.find(qn("a:buAutoNum")) is not None:
        return "ol"
    if ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buBlip")) is not None:
        return "ul"
    return "inherit"


def find_placeholder(placeholders, idx=None, ph_type=None):
    for ph in placeholders:
        pf = ph.placeholder_format
        if idx is not None and pf.idx == idx:
            return ph
        if ph_type is not None and pf.type == ph_type:
            return ph
    return None


def list_style_of(shape_element):
    tx_body = shape_element.find(qn("p:txBody"))
    if tx_body is None:
        return None
    return tx_body.find(qn("a:lstStyle"))


def style_chain(shape, slide):
    """Yield the level-style containers a paragraph inherits from, nearest first."""
    yield list_style_of(shape._element)

    layout = slide.slide_layout
    master = layout.slide_master
    ph_type = None
    if shape.is_placeholder:
        pf = shape.placeholder_format
        ph_type = pf.type
        base_type = ph_type
        layout_ph = find_placeholder(layout.placeholders, idx=pf.idx)
        if layout_ph is not None:
            yield list_style_of(layout_ph._element)
            base_type = layout_ph.placeholder_format.type
        master_ph = find_placeholder(master.placeholders, ph_type=base_type)
        if master_ph is None and base_type in BODY_LIKE_PLACEHOLDERS:
            master_ph = find_placeholder(master.placeholders, ph_type=PP_PLACEHOLDER.BODY)
        if master_ph is None and base_type in TITLE_PLACEHOLDERS:
            master_ph = find_placeholder(master.placeholders, ph_type=PP_PLACEHOLDER.TITLE)
        if master_ph is not None:
            yield list_style_of(master_ph._element)

    tx_styles = master._element.find(qn("p:txStyles"))
    if tx_styles is not None:
        if ph_type in TITLE_PLACEHOLDERS:
            yield tx_styles.find(qn("p:titleStyle"))
        elif ph_type in BODY_LIKE_PLACEHOLDERS:
            yield tx_styles.find(qn("p:bodyStyle"))
        else:
            yield tx_styles.find(qn("p:otherStyle"))


def bullet_kind(paragraph, shape, slide):
    kind = bullet_from_ppr(paragraph._p.find(qn("a:pPr")))
    if kind != "inherit":
        return kind
    level = paragraph.level
    for container in style_chain(shape, slide):
        if container is None:
            continue
        kind = bullet_from_ppr(container.find(qn(f"a:lvl{level + 1}pPr")))
        if kind != "inherit":
            return kind
    return None


# --------------------------------------------------------------------------- #
# Inline text
# --------------------------------------------------------------------------- #

def render_inline(paragraph, part, report: Report) -> str | None:
    """Return the paragraph's inner HTML, or None if it has no visible text.

    Slide-number fields are dropped. Adjacent runs that share a hyperlink are
    merged into one anchor. Line breaks become <br>.
    """
    pieces: list[tuple[str | None, str]] = []   # (href, text) or ("<br>", "")
    for child in paragraph._p:
        tag = child.tag
        if tag == qn("a:r"):
            text = child.findtext(qn("a:t")) or ""
            if not text:
                continue
            rpr = child.find(qn("a:rPr"))
            href = resolve_link(part, rpr.find(qn("a:hlinkClick"))) if rpr is not None else None
            pieces.append((href, text))
        elif tag == qn("a:fld"):
            if child.get("type") == "slidenum":
                continue
            text = child.findtext(qn("a:t")) or ""
            if text:
                pieces.append((None, text))
        elif tag == qn("a:br"):
            pieces.append(("<br>", ""))

    if not any(text.strip() for href, text in pieces if href != "<br>"):
        return None

    # Trim leading and trailing breaks.
    while pieces and pieces[0][0] == "<br>":
        pieces.pop(0)
    while pieces and pieces[-1][0] == "<br>":
        pieces.pop()

    out: list[str] = []
    open_href: str | None = None
    for href, text in pieces:
        if href == "<br>":
            if open_href is not None:
                out.append("</a>")
                open_href = None
            out.append("<br>")
            continue
        if href != open_href:
            if open_href is not None:
                out.append("</a>")
            if href is not None:
                out.append(f'<a href="{attr(href)}">')
                report.links.add(href)
            open_href = href
        out.append(esc(text))
    if open_href is not None:
        out.append("</a>")
    return "".join(out)


def plain_text(paragraphs) -> str:
    text = " ".join(p.text.replace("\v", " ").replace("\n", " ") for p in paragraphs)
    return re.sub(r"\s+", " ", text).strip()


# --------------------------------------------------------------------------- #
# Shapes
# --------------------------------------------------------------------------- #

def ordered_shapes(shapes, exclude=None):
    """Reading order: rows top to bottom, shapes within a row left to right."""
    # Compare underlying XML elements: python-pptx hands out a new proxy object
    # on every access, so identity on the proxies would never match.
    exclude_element = exclude._element if exclude is not None else None
    items = [s for s in shapes if s._element is not exclude_element]
    items.sort(key=lambda s: ((s.top or 0), (s.left or 0)))
    rows: list[list] = []
    current: list = []
    row_top = 0
    for shape in items:
        top = shape.top or 0
        if current and top - row_top > ROW_TOLERANCE:
            rows.append(current)
            current = []
        if not current:
            row_top = top
        current.append(shape)
    if current:
        rows.append(current)
    ordered = []
    for row in rows:
        row.sort(key=lambda s: ((s.left or 0), (s.top or 0)))
        ordered.extend(row)
    return ordered


def render_text_shape(shape, slide, report: Report, slide_height: int,
                      label: str | None = None) -> list[str]:
    """Render a text shape. A label (the alt text of an icon that sits beside
    the shape, such as "Instagram:") is written in front of the first
    paragraph, outside any hyperlink, so the line reads "Instagram: @UM_NaNDA".
    """
    if not shape.has_text_frame:
        return []
    paragraphs = list(shape.text_frame.paragraphs)
    text = plain_text(paragraphs)
    if not text:
        return []
    # A bare number sitting at the bottom edge is a slide number, not content.
    if re.fullmatch(r"\d{1,3}", text) and (shape.top or 0) > slide_height * 0.9:
        return []

    link = shape_link(shape)
    has_run_links = any(
        r.find(qn("a:rPr")) is not None
        and r.find(qn("a:rPr")).find(qn("a:hlinkClick")) is not None
        for p in paragraphs for r in p._p.findall(qn("a:r"))
    )
    wrap_link = link if (link and not has_run_links) else None

    out: list[str] = []
    stack: list[tuple[int, str]] = []   # open lists as (level, tag)

    def close_all():
        while stack:
            out.append(f"</li></{stack.pop()[1]}>")

    for para in paragraphs:
        inner = render_inline(para, shape.part, report)
        if inner is None:
            continue
        if wrap_link:
            inner = f'<a href="{attr(wrap_link)}">{inner}</a>'
            report.links.add(wrap_link)
            wrap_link = None
        if label:
            inner = f"{esc(label)} {inner}"
            label = None
        kind = bullet_kind(para, shape, slide)
        if kind is None:
            close_all()
            out.append(f"<p>{inner}</p>")
            continue
        level = para.level
        while stack and stack[-1][0] > level:
            out.append(f"</li></{stack.pop()[1]}>")
        if stack and stack[-1][0] == level:
            if stack[-1][1] == kind:
                out.append("</li>")
            else:
                out.append(f"</li></{stack.pop()[1]}>")
        if not stack or stack[-1][0] < level:
            out.append(f"<{kind}>")
            stack.append((level, kind))
        out.append(f"<li>{inner}")
    close_all()
    return out


def render_table(shape, report: Report) -> list[str]:
    table = shape.table
    rows = list(table.rows)
    if not rows:
        return []
    out = ["<table>"]
    for r_index, row in enumerate(rows):
        cells_html = []
        for cell in row.cells:
            if cell.is_spanned:
                continue
            span = ""
            if cell.is_merge_origin:
                if cell.span_width > 1:
                    span += f' colspan="{cell.span_width}"'
                if cell.span_height > 1:
                    span += f' rowspan="{cell.span_height}"'
            parts = [render_inline(p, shape.part, report) for p in cell.text_frame.paragraphs]
            inner = "<br>".join(p for p in parts if p)
            if r_index == 0:
                cells_html.append(f'<th scope="col"{span}>{inner}</th>')
            else:
                cells_html.append(f"<td{span}>{inner}</td>")
        row_html = "<tr>" + "".join(cells_html) + "</tr>"
        if r_index == 0:
            out.append(f"<thead>{row_html}</thead><tbody>")
        else:
            out.append(row_html)
    out.append("</tbody></table>")
    return out


def render_described(shape, kind: str, slide_no: int, report: Report) -> list[str]:
    """Render a picture, chart, diagram, or other object as its alt text."""
    if is_decorative(shape):
        return []
    alt = alt_text(shape)
    if not alt or FILENAME_ALT.match(alt) or alt.lower() == "preencoded.png":
        report.missing_alt.append((slide_no, kind, shape.name))
        return []
    text = f"{kind}: {esc(alt)}"
    link = shape_link(shape)
    if link:
        report.links.add(link)
        return [f'<p class="image"><a href="{attr(link)}">{text}</a></p>']
    return [f'<p class="image">{text}</p>']


def is_label_picture(shape) -> bool:
    """A non-decorative picture whose alt text ends with a colon."""
    return (
        shape._element.tag == qn("p:pic")
        and not is_decorative(shape)
        and alt_text(shape).endswith(":")
    )


def has_visible_text(shape) -> bool:
    return (
        shape._element.tag == qn("p:sp")
        and shape.has_text_frame
        and bool(plain_text(shape.text_frame.paragraphs))
    )


def render_shapes(shapes, slide, slide_no: int, report: Report, slide_height: int,
                  exclude=None) -> list[str]:
    """Render a collection of shapes in reading order.

    Before rendering, every label picture (see is_label_picture) is paired with
    the nearest text shape to its right whose top edge is within ROW_TOLERANCE
    of the picture's. The picture is then dropped and its alt text is written
    in front of that shape's first paragraph. A label with no partner is
    rendered as an ordinary image description.
    """
    ordered = ordered_shapes(shapes, exclude=exclude)
    labels: dict[int, str] = {}      # id(text shape element) -> label
    consumed: set[int] = set()       # id(label picture element)
    for pic in ordered:
        if not is_label_picture(pic):
            continue
        partner = None
        for candidate in ordered:
            if candidate._element is pic._element or not has_visible_text(candidate):
                continue
            if id(candidate._element) in labels:
                continue
            if (candidate.left or 0) < (pic.left or 0):
                continue
            if abs((candidate.top or 0) - (pic.top or 0)) > ROW_TOLERANCE:
                continue
            if partner is None or (candidate.left or 0) < (partner.left or 0):
                partner = candidate
        if partner is not None:
            labels[id(partner._element)] = alt_text(pic)
            consumed.add(id(pic._element))

    out: list[str] = []
    for shape in ordered:
        if id(shape._element) in consumed:
            continue
        out.extend(render_shape(shape, slide, slide_no, report, slide_height,
                                label=labels.get(id(shape._element))))
    return out


def render_shape(shape, slide, slide_no: int, report: Report, slide_height: int,
                 label: str | None = None) -> list[str]:
    element = shape._element
    tag = element.tag

    if shape.is_placeholder:
        try:
            if shape.placeholder_format.type in SKIP_PLACEHOLDERS:
                return []
        except ValueError:
            pass

    if tag == qn("p:grpSp"):
        return render_shapes(shape.shapes, slide, slide_no, report, slide_height)

    if tag == qn("p:pic"):
        nv_pr = element.find(qn("p:nvPicPr"))
        kind = "Image"
        if nv_pr is not None:
            inner_nv = nv_pr.find(qn("p:nvPr"))
            if inner_nv is not None:
                if inner_nv.find(qn("a:videoFile")) is not None:
                    kind = "Video"
                elif inner_nv.find(qn("a:audioFile")) is not None:
                    kind = "Audio"
        return render_described(shape, kind, slide_no, report)

    if tag == qn("p:graphicFrame"):
        if getattr(shape, "has_table", False) and shape.has_table:
            return render_table(shape, report)
        if getattr(shape, "has_chart", False) and shape.has_chart:
            return render_described(shape, "Chart", slide_no, report)
        graphic_data = element.find(qn("a:graphic") + "/" + qn("a:graphicData"))
        uri = graphic_data.get("uri", "") if graphic_data is not None else ""
        kind = "Diagram" if uri.endswith("/diagram") else "Object"
        return render_described(shape, kind, slide_no, report)

    if tag == qn("p:cxnSp"):
        return []

    return render_text_shape(shape, slide, report, slide_height, label=label)


# --------------------------------------------------------------------------- #
# Page assembly
# --------------------------------------------------------------------------- #

def slide_title(slide) -> str:
    title_shape = slide.shapes.title
    if title_shape is None or not title_shape.has_text_frame:
        return ""
    return plain_text(title_shape.text_frame.paragraphs)


def build_page(stem: str, deck_title: str, presenters: str, heading: str | None,
               prs, report: Report) -> str:
    slide_height = prs.slide_height or 1
    page_heading = heading or deck_title
    sections: list[str] = []
    nav_items: list[str] = []

    for n, slide in enumerate(prs.slides, start=1):
        title = slide_title(slide)
        if title:
            heading = f"Slide {n}: {esc(title)}"
        else:
            heading = f"Slide {n}"
            report.missing_title.append(n)
        nav_items.append(f'<li><a href="#slide-{n}">{heading}</a></li>')

        body = render_shapes(slide.shapes, slide, n, report, slide_height,
                             exclude=slide.shapes.title)

        sections.append(
            f'<section id="slide-{n}" aria-labelledby="slide-{n}-heading">\n'
            f'<h2 id="slide-{n}-heading">{heading}</h2>\n'
            + "\n".join(body)
            + "\n</section>"
        )

    report.slide_count = len(prs.slides)
    page_title = f"{deck_title} (slide transcript)"

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(page_title)}</title>
<style>
{CSS}
</style>
</head>
<body>
<a class="skip-link" href="#main">Skip to slide content</a>
<div class="page">
<header>
<h1>{esc(page_heading)}</h1>
<p class="about">{esc(WORKSHOP_TITLE)}. {esc(EVENT_LINE)}. Presented by {esc(presenters)}, {esc(DATE_LINE)}.</p>
<p>This page is a text transcript of the slides, one section per slide, with image descriptions written out where the slides carry images. You can also <a href="{attr(stem)}.pptx">download this deck as a PowerPoint file</a>, <a href="{attr(stem)}.pdf">download this deck as a PDF</a>, or <a href="{attr(REPO_URL)}">return to the workshop materials repository on GitHub</a>.</p>
</header>
<nav aria-label="Slides">
<h2 id="contents">Slides in this deck</h2>
<ol>
{chr(10).join(nav_items)}
</ol>
</nav>
<main id="main">
{chr(10).join(sections)}
</main>
</div>
</body>
</html>
"""


def convert(stem: str, deck_title: str, presenters: str, heading: str | None) -> Report:
    source = SLIDES_DIR / f"{stem}.pptx"
    target = SLIDES_DIR / f"{stem}.html"
    if not source.exists():
        raise SystemExit(f"Missing public deck: {source}. Run tools\\make_public_decks.ps1 first.")
    report = Report(stem=stem)
    prs = Presentation(str(source))
    page = build_page(stem, deck_title, presenters, heading, prs, report)
    target.write_text(page, encoding="utf-8", newline="\n")
    return report


def main(argv: list[str]) -> int:
    wanted = set(argv[1:])
    decks = [d for d in DECKS if not wanted or d[0] in wanted]
    unknown = wanted - {d[0] for d in DECKS}
    if unknown:
        print("Unknown deck stem(s):", ", ".join(sorted(unknown)))
        return 2

    for stem, deck_title, presenters, heading in decks:
        report = convert(stem, deck_title, presenters, heading)
        print(f"{stem}.html: {report.slide_count} slides, {len(report.links)} distinct links")
        for slide_no, kind, name in report.missing_alt:
            print(f"   missing alt text: slide {slide_no}, {kind.lower()} {name!r}")
        for slide_no in report.missing_title:
            print(f"   no title placeholder: slide {slide_no} (heading is 'Slide {slide_no}')")
        for link in sorted(report.links):
            print(f"   link: {link}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
