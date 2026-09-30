"""Prove the public deck copies are clean, without relying on PowerPoint.

For each public deck in slides/<stem>.pptx the script reads the file as a ZIP
and checks that:

* there are no comment parts (ppt/comments/*, ppt/commentAuthors*), no comment
  or author relationships, and no comment content types;
* there is no speaker-notes text: every <a:t> in ppt/notesSlides/*.xml is
  empty, a slide number, or the slide-number field marker;
* there are no hidden slides (<p:sld show="0">);
* the slide count matches EXPECTED_SLIDES, both from the ZIP and from
  python-pptx opening the file.

For each slides/<stem>.pdf it checks, with PyMuPDF, that:

* the page count equals the public slide count;
* text is selectable (it prints the first line of text on page 4);
* the file is tagged (/MarkInfo with /Marked true, and a /StructTreeRoot).

For each slides/<stem>.html that exists it checks that:

* the number of slide headings (h2 with an id ending in "-heading") equals the
  public slide count and none of them is empty;
* the page has lang="en", a <title>, a skip link to #main, and a <nav>;
* every relative link resolves to a file in the repository;
* every external hyperlink in the public .pptx appears in the HTML.

Exit status is 0 when every check passes and 1 otherwise.

Usage:
    python tools/check_public_decks.py

Requires python-pptx and pymupdf (pip install python-pptx pymupdf).
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

REPO_ROOT = Path(__file__).resolve().parents[1]
SLIDES_DIR = REPO_ROOT / "slides"

# Public slide counts: original count minus hidden slides.
EXPECTED_SLIDES = {
    "01_Welcome_and_Overview": 28,
    "02_Neighborhood_Environment_and_Aging_with_Disability": 16,
    "03_NaNDA_and_Immune_Aging": 19,
    "04_Choosing_Your_First_NaNDA_Dataset": 34,
    "05_Reconvene_and_Wrap-Up": 9,
}

NOTES_TEXT_ALLOWED = re.compile(r"^\s*(\d+|‹#›)?\s*$")
COMMENT_TYPE_MARKERS = ("comments+xml", "commentauthors+xml", "authors+xml")

failures: list[str] = []


def fail(message: str) -> None:
    failures.append(message)
    print(f"   FAIL: {message}")


def ok(message: str) -> None:
    print(f"   ok:   {message}")


def check_pptx(stem: str, expected: int) -> int:
    path = SLIDES_DIR / f"{stem}.pptx"
    print(f"{path.name}")
    if not path.exists():
        fail("file is missing")
        return 0
    with zipfile.ZipFile(path) as z:
        names = z.namelist()

        comment_parts = [
            n for n in names
            if n.startswith("ppt/comments/") or n in ("ppt/authors.xml", "ppt/commentAuthors.xml")
        ]
        content_types = z.read("[Content_Types].xml").decode("utf-8", "ignore").lower()
        comment_types = [m for m in COMMENT_TYPE_MARKERS if m in content_types]
        comment_rels = []
        for n in names:
            if n.endswith(".rels"):
                rels = z.read(n).decode("utf-8", "ignore").lower()
                if "/comments" in rels or "/commentauthors" in rels or "relationships/authors" in rels:
                    comment_rels.append(n)
        if comment_parts or comment_types or comment_rels:
            fail(f"comments remain: parts={comment_parts} types={comment_types} rels={comment_rels}")
        else:
            ok("no comment parts, relationships, or content types")

        leaks = []
        notes_parts = [n for n in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n)]
        for n in notes_parts:
            xml = z.read(n).decode("utf-8", "ignore")
            for text in re.findall(r"<a:t>([^<]*)</a:t>", xml):
                if not NOTES_TEXT_ALLOWED.match(text):
                    leaks.append((n, text[:60]))
        if leaks:
            fail(f"speaker-notes text remains in {len(leaks)} place(s), for example {leaks[0]}")
        else:
            ok(f"no speaker-notes text in {len(notes_parts)} notes parts")

        hidden = []
        slide_parts = [n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)]
        for n in slide_parts:
            xml = z.read(n).decode("utf-8", "ignore")
            if re.search(r'<p:sld\b[^>]*\bshow="0"', xml):
                hidden.append(n)
        if hidden:
            fail(f"hidden slides remain: {hidden}")
        else:
            ok("no hidden slides")

    zip_count = len(slide_parts)
    pptx_count = len(Presentation(str(path)).slides)
    if zip_count == pptx_count == expected:
        ok(f"slide count {expected}")
    else:
        fail(f"slide count: zip {zip_count}, python-pptx {pptx_count}, expected {expected}")
    return pptx_count


def check_pdf(stem: str, expected: int) -> None:
    path = SLIDES_DIR / f"{stem}.pdf"
    print(f"{path.name}")
    if not path.exists():
        fail("file is missing")
        return
    try:
        import fitz  # PyMuPDF
    except ImportError:
        fail("pymupdf is not installed (pip install pymupdf), PDF checks skipped")
        return
    doc = fitz.open(str(path))
    if doc.page_count == expected:
        ok(f"page count {expected}")
    else:
        fail(f"page count {doc.page_count}, expected {expected}")

    page_index = min(3, doc.page_count - 1)
    text = doc[page_index].get_text().strip()
    first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")
    if first_line:
        ok(f"selectable text on page {page_index + 1}: {first_line[:70]!r}")
    else:
        fail(f"no selectable text on page {page_index + 1}")

    catalog = doc.pdf_catalog()
    marked_type, marked_value = doc.xref_get_key(catalog, "MarkInfo/Marked")
    struct_type, _ = doc.xref_get_key(catalog, "StructTreeRoot")
    if marked_value == "true" and struct_type != "null":
        ok("tagged PDF (/MarkInfo /Marked true, /StructTreeRoot present)")
    else:
        fail(f"not tagged: Marked={marked_type}:{marked_value}, StructTreeRoot={struct_type}")
    title = (doc.metadata or {}).get("title", "")
    print(f"   info: PDF title metadata {title!r}")
    doc.close()


def pptx_external_links(path: Path) -> set[str]:
    links: set[str] = set()
    prs = Presentation(str(path))
    for slide in prs.slides:
        part = slide.part
        for hlink in slide._element.iter(qn("a:hlinkClick")):
            rid = hlink.get(qn("r:id"))
            if not rid:
                continue
            try:
                rel = part.rels[rid]
            except KeyError:
                continue
            if rel.is_external:
                links.add(rel.target_ref)
    return links


def check_html(stem: str, expected: int) -> None:
    path = SLIDES_DIR / f"{stem}.html"
    if not path.exists():
        print(f"{path.name}: not present, transcript checks skipped")
        return
    print(f"{path.name}")
    text = path.read_text(encoding="utf-8")

    headings = re.findall(r'<h2 id="slide-\d+-heading">(.*?)</h2>', text)
    if len(headings) == expected:
        ok(f"{expected} slide headings")
    else:
        fail(f"{len(headings)} slide headings, expected {expected}")
    empty = [h for h in headings if not re.sub(r"<[^>]+>", "", h).strip()]
    if empty:
        fail(f"{len(empty)} empty slide heading(s)")

    for label, pattern in [
        ('lang="en"', r'<html lang="en">'),
        ("title", r"<title>[^<]+</title>"),
        ("skip link", r'<a class="skip-link" href="#main">'),
        ("nav", r'<nav aria-label="Slides">'),
        ("main", r'<main id="main">'),
    ]:
        if re.search(pattern, text):
            ok(f"{label} present")
        else:
            fail(f"{label} missing")

    hrefs = re.findall(r'href="([^"]+)"', text)
    unresolved = []
    for href in hrefs:
        if href.startswith(("#", "http://", "https://", "mailto:")):
            continue
        if not (path.parent / href).exists():
            unresolved.append(href)
    if unresolved:
        fail(f"relative links that do not resolve: {sorted(set(unresolved))}")
    else:
        ok("every relative link resolves to a file")

    pptx_links = pptx_external_links(SLIDES_DIR / f"{stem}.pptx")
    missing = sorted(link for link in pptx_links if f'href="{link}"' not in text and link.replace("&", "&amp;") not in text)
    if missing:
        fail(f"hyperlinks in the .pptx that are not in the HTML: {missing}")
    else:
        ok(f"all {len(pptx_links)} hyperlinks from the .pptx appear in the HTML")


def main() -> int:
    for stem, expected in EXPECTED_SLIDES.items():
        check_pptx(stem, expected)
        check_pdf(stem, expected)
        check_html(stem, expected)
        print()
    if failures:
        print(f"{len(failures)} check(s) failed.")
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
